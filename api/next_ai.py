"""Session-scoped hybrid RAG and Graph-RAG chatbot for uploaded evidence."""
from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import os
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Optional

from loguru import logger
from pydantic import BaseModel, Field


class NextAIContextRequest(BaseModel):
    tracking_id: Optional[str] = None
    industry: str = ""
    use_case: str = ""
    objective: str = ""
    persona: str = ""
    hint_topics: list[str] = Field(default_factory=list)
    primary_entities: list[str] = Field(default_factory=list)
    key_relationships: list[str] = Field(default_factory=list)
    tabular_files: list[str] = Field(default_factory=list)
    document_files: list[str] = Field(default_factory=list)
    image_files: list[str] = Field(default_factory=list)
    audio_files: list[str] = Field(default_factory=list)
    text_files: list[str] = Field(default_factory=list)


class NextAIAskRequest(NextAIContextRequest):
    question: str
    history: list[dict[str, str]] = Field(default_factory=list)
    allow_llm_synthesis: bool = False


_INDEX_CACHE: dict[str, dict[str, Any]] = {}
_CHROMA_COLLECTION = None

_TOKEN_RE = re.compile(r"[a-zA-Z][a-zA-Z0-9_-]{1,}")
_SENSITIVE_HEADER = re.compile(
    r"name|email|phone|mobile|address|dob|birth|ssn|aadhaar|passport|pan|"
    r"account.?number|iban|card.?number",
    re.IGNORECASE,
)
_STOPWORDS = {
    "the", "and", "for", "that", "this", "with", "from", "into", "have", "has",
    "are", "was", "were", "will", "would", "could", "should", "about", "what",
    "which", "when", "where", "who", "how", "your", "their", "there", "than",
    "then", "data", "file", "record", "document",
    "pii", "field", "fields", "column", "columns", "value", "values", "source",
}
_SEMANTIC_EXPANSIONS = {
    "fraud": ("risk", "suspicious", "anomaly", "investigation"),
    "laundering": ("aml", "transaction", "suspicious", "money"),
    "customer": ("client", "subscriber", "patient", "buyer"),
    "supplier": ("vendor", "procurement", "source"),
    "fault": ("alarm", "incident", "outage", "failure"),
    "risk": ("severity", "exposure", "flag", "alert"),
    "relationship": ("link", "connected", "edge", "association"),
    "summary": ("overview", "key", "main", "important"),
    "obligation": ("must", "shall", "required", "compliance"),
}
_BLOCKED_PATTERNS = (
    "ignore previous instructions", "reveal system prompt", "show api key",
    "dump credentials", "bypass guardrail", "delete database", "drop database",
    "disable security", "exfiltrate", "steal password",
)


def _clean(value: Any, limit: int = 2000) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()[:limit]


def _mask_pii(text: str) -> tuple[str, int]:
    count = 0
    patterns = [
        (
            re.compile(
                r"(?i)\b(full name|customer name|individual name|person name|"
                r"beneficiary name|residential address)\s*[:=]\s*"
                r"[A-Z0-9][A-Z0-9 .,'/#-]{2,80}"
            ),
            "[PII_FIELD]=[PII]",
        ),
        (re.compile(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", re.I), "[EMAIL]"),
        (re.compile(r"(?<!\w)(?:\+?\d[\d .()-]{8,}\d)(?!\w)"), "[PHONE/ID]"),
        (re.compile(r"\b(?:\d[ -]*?){13,19}\b"), "[CARD/ACCOUNT]"),
        (re.compile(r"\b\d{3}-\d{2}-\d{4}\b"), "[GOVERNMENT_ID]"),
    ]
    masked = text
    for pattern, replacement in patterns:
        masked, n = pattern.subn(replacement, masked)
        count += n
    return masked, count


def _tabular_text(data: bytes, delimiter: str = ",") -> tuple[str, int]:
    decoded = data.decode("utf-8-sig", errors="replace")
    rows = list(csv.reader(io.StringIO(decoded), delimiter=delimiter))
    if not rows:
        return "", 0
    headers = [str(item or f"column_{i + 1}") for i, item in enumerate(rows[0])]
    lines = ["Columns: " + ", ".join(headers)]
    pii_count = 0
    for row in rows[1:301]:
        fields = []
        for index, header in enumerate(headers[:40]):
            value = str(row[index] if index < len(row) else "")
            if value and _SENSITIVE_HEADER.search(header):
                value = "[PII]"
                pii_count += 1
            else:
                value, n = _mask_pii(value)
                pii_count += n
            fields.append(f"{header}={value}")
        lines.append(" | ".join(fields))
    return "\n".join(lines), pii_count


def _extract_text(filename: str, data: bytes) -> tuple[str, int]:
    ext = Path(filename).suffix.lower()
    try:
        if ext in (".csv", ".tsv"):
            return _tabular_text(data, "\t" if ext == ".tsv" else ",")
        if ext in (".xlsx", ".xls"):
            import openpyxl

            book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            sheet = book.active
            rows = list(sheet.iter_rows(values_only=True, max_row=301, max_col=40))
            if not rows:
                return "", 0
            buffer = io.StringIO()
            csv.writer(buffer).writerows(rows)
            return _tabular_text(buffer.getvalue().encode("utf-8"))
        if ext == ".pdf":
            try:
                from pypdf import PdfReader

                reader = PdfReader(io.BytesIO(data))
                text = "\n".join((page.extract_text() or "") for page in reader.pages[:30])
            except Exception:
                import pdfplumber

                with pdfplumber.open(io.BytesIO(data)) as pdf:
                    text = "\n".join((page.extract_text() or "") for page in pdf.pages[:30])
        elif ext in (".doc", ".docx"):
            import docx

            document = docx.Document(io.BytesIO(data))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        elif ext in (".txt", ".md", ".rtf", ".json", ".xml", ".log"):
            text = data[:500_000].decode("utf-8", errors="replace")
        elif ext in (".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff"):
            from PIL import Image
            import pytesseract

            text = pytesseract.image_to_string(Image.open(io.BytesIO(data)))
        elif ext in (".wav", ".aiff", ".flac"):
            import speech_recognition as sr

            recognizer = sr.Recognizer()
            with sr.AudioFile(io.BytesIO(data)) as source:
                audio = recognizer.record(source)
            text = recognizer.recognize_sphinx(audio)
        else:
            return "", 0
        return _mask_pii(text)
    except Exception as exc:
        logger.warning(f"Next AI could not extract {filename}: {exc}")
        return "", 0


def _chunk_text(text: str, size: int = 900, overlap: int = 140) -> list[str]:
    text = re.sub(r"\r\n?", "\n", text).strip()
    if not text:
        return []
    chunks = []
    start = 0
    while start < len(text) and len(chunks) < 500:
        end = min(len(text), start + size)
        if end < len(text):
            split = max(text.rfind("\n", start, end), text.rfind(". ", start, end))
            if split > start + size // 2:
                end = split + 1
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(text):
            break
        start = max(start + 1, end - overlap)
    return chunks


def _tokens(text: str) -> list[str]:
    base = []
    for token in _TOKEN_RE.findall(text):
        normalized = token.lower()
        if len(normalized) > 4 and normalized.endswith("s") and not normalized.endswith("ss"):
            normalized = normalized[:-1]
        if len(normalized) > 2 and normalized not in _STOPWORDS:
            base.append(normalized)
    expanded = list(base)
    for token in base:
        expanded.extend(_SEMANTIC_EXPANSIONS.get(token, ()))
    return expanded


def _embedding(text: str, dimensions: int = 384) -> list[float]:
    vector = [0.0] * dimensions
    counts = Counter(_tokens(text))
    for token, count in counts.items():
        digest = hashlib.blake2b(token.encode("utf-8"), digest_size=8).digest()
        index = int.from_bytes(digest, "big") % dimensions
        vector[index] += 1.0 + math.log(count)
    norm = math.sqrt(sum(value * value for value in vector)) or 1.0
    return [value / norm for value in vector]


def _cosine(left: list[float], right: list[float]) -> float:
    return sum(a * b for a, b in zip(left, right))


def _entities(text: str) -> list[str]:
    identifiers = re.findall(r"\b[A-Z]{2,12}[-_/]\d{2,16}\b", text)
    phrases = re.findall(r"\b[A-Z][A-Za-z0-9&.-]+(?:\s+[A-Z][A-Za-z0-9&.-]+){0,2}\b", text)
    blocked = {"Columns", "PII", "Email", "Phone", "Reference"}
    return list(dict.fromkeys(
        value[:100] for value in identifiers + phrases
        if value not in blocked and not value.startswith("[")
    ))[:20]


def _request_filenames(req: NextAIContextRequest) -> list[str]:
    return list(dict.fromkeys(
        req.tabular_files + req.document_files + req.image_files
        + req.audio_files + req.text_files
    ))


def _source_modality(filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext in (".csv", ".tsv", ".xlsx", ".xls"):
        return "structured"
    if ext in (".png", ".jpg", ".jpeg", ".webp"):
        return "image"
    if ext in (".mp3", ".wav", ".m4a", ".aac"):
        return "audio"
    return "unstructured"


def _collect_uploaded_sources(req: NextAIContextRequest, db_client=None) -> list[dict]:
    sources: dict[str, dict] = {
        filename: {
            "filename": filename,
            "modality": _source_modality(filename),
            "text": "",
            "pii_masked": 0,
        }
        for filename in _request_filenames(req)
    }
    if req.tracking_id and db_client:
        for item in db_client.get_files_for_session(req.tracking_id):
            content = db_client.get_file_content(item["id"]) or {}
            filename = item.get("filename") or content.get("filename") or ""
            if not filename:
                continue
            text, pii_count = _extract_text(filename, content.get("file_data") or b"")
            sources[filename] = {
                "file_id": item.get("id"),
                "filename": filename,
                "modality": _source_modality(filename),
                "text": text,
                "pii_masked": pii_count,
            }
    return list(sources.values())


def _cache_key(req: NextAIContextRequest, sources: list[dict]) -> str:
    signature = "|".join(
        f"{source.get('file_id', '')}:{source['filename']}:{len(source.get('text') or '')}"
        for source in sources
    )
    return hashlib.sha1(
        f"{req.tracking_id or 'no-session'}|{signature}".encode("utf-8")
    ).hexdigest()


def _get_chroma_collection():
    global _CHROMA_COLLECTION
    if _CHROMA_COLLECTION is not None:
        return _CHROMA_COLLECTION
    try:
        import chromadb

        path = os.path.join(os.path.dirname(__file__), "..", "data", "next_ai_chroma")
        client = chromadb.PersistentClient(path=os.path.abspath(path))
        _CHROMA_COLLECTION = client.get_or_create_collection(
            name="next_ai_uploaded_evidence",
            metadata={"hnsw:space": "cosine"},
        )
    except Exception as exc:
        logger.warning(f"Next AI ChromaDB unavailable; using in-memory vectors: {exc}")
        _CHROMA_COLLECTION = False
    return _CHROMA_COLLECTION


def _build_index(req: NextAIContextRequest, db_client=None) -> dict:
    sources = _collect_uploaded_sources(req, db_client)
    key = _cache_key(req, sources)
    if key in _INDEX_CACHE:
        return _INDEX_CACHE[key]

    chunks: list[dict] = []
    entity_sources: dict[str, set[str]] = defaultdict(set)
    pii_count = 0
    for source in sources:
        pii_count += int(source.get("pii_masked") or 0)
        for index, text in enumerate(_chunk_text(source.get("text") or "")):
            chunk_id = hashlib.sha1(
                f"{key}|{source['filename']}|{index}|{text[:80]}".encode("utf-8")
            ).hexdigest()
            entities = _entities(text)
            for entity in entities:
                entity_sources[entity].add(source["filename"])
            chunks.append({
                "id": chunk_id,
                "filename": source["filename"],
                "modality": source["modality"],
                "chunk_index": index,
                "text": text,
                "entities": entities,
                "embedding": _embedding(text),
            })

    collection = _get_chroma_collection()
    vector_store = (
        "ChromaDB cosine index"
        if collection is not False and collection is not None
        else "in-memory cosine"
    )
    if collection is not False and collection is not None and chunks:
        try:
            collection.upsert(
                ids=[chunk["id"] for chunk in chunks],
                embeddings=[chunk["embedding"] for chunk in chunks],
                documents=[chunk["text"] for chunk in chunks],
                metadatas=[{
                    "tracking_key": key,
                    "filename": chunk["filename"],
                    "modality": chunk["modality"],
                    "chunk_index": chunk["chunk_index"],
                } for chunk in chunks],
            )
        except Exception as exc:
            logger.warning(f"Next AI Chroma indexing failed: {exc}")
            vector_store = "in-memory cosine"

    index = {
        "key": key,
        "tracking_id": req.tracking_id,
        "sources": sources,
        "chunks": chunks,
        "entity_sources": entity_sources,
        "pii_masked": pii_count,
        "vector_store": vector_store,
    }
    _INDEX_CACHE[key] = index
    return index


def _semantic_search(index: dict, question: str, limit: int = 6) -> list[dict]:
    if not index["chunks"]:
        return []
    query_vector = _embedding(question)
    query_terms = set(_tokens(question))
    ranked = []
    collection = _get_chroma_collection()
    if collection is not False and collection is not None:
        try:
            result = collection.query(
                query_embeddings=[query_vector],
                n_results=min(max(limit * 4, 20), len(index["chunks"])),
                where={"tracking_key": index["key"]},
                include=["documents", "metadatas", "distances"],
            )
            for document, metadata, distance in zip(
                result["documents"][0], result["metadatas"][0], result["distances"][0]
            ):
                ranked.append({
                    "filename": metadata["filename"],
                    "modality": metadata["modality"],
                    "chunk_index": metadata["chunk_index"],
                    "text": document,
                    "score": round(max(0.0, 1.0 - float(distance)), 4),
                    "term_overlap": len(query_terms & set(_tokens(document))),
                })
        except Exception as exc:
            logger.warning(f"Next AI Chroma search failed: {exc}")

    for chunk in index["chunks"]:
        ranked.append({
            "filename": chunk["filename"],
            "modality": chunk["modality"],
            "chunk_index": chunk["chunk_index"],
            "text": chunk["text"],
            "score": round(_cosine(query_vector, chunk["embedding"]), 4),
            "term_overlap": len(query_terms & set(_tokens(chunk["text"]))),
        })
    deduplicated = {}
    for item in ranked:
        key = (item["filename"], item["chunk_index"])
        previous = deduplicated.get(key)
        if previous is None or item["score"] > previous["score"]:
            deduplicated[key] = item
    return sorted(
        deduplicated.values(),
        key=lambda item: (
            item.get("term_overlap", 0) * 0.12 + item["score"]
        ),
        reverse=True,
    )[:limit]


def _local_graph_context(index: dict, evidence: list[dict], question: str) -> dict:
    evidence_files = {item["filename"] for item in evidence}
    query_tokens = set(_tokens(question))
    entities = []
    for entity, filenames in index["entity_sources"].items():
        overlap = bool(query_tokens & set(_tokens(entity)))
        if overlap or filenames & evidence_files:
            entities.append((entity, sorted(filenames)))
    entities = entities[:18]
    nodes = [
        {"id": f"source:{filename}", "label": filename, "type": "Source"}
        for filename in sorted(evidence_files)
    ]
    edges = []
    for entity, filenames in entities:
        entity_id = "entity:" + hashlib.md5(entity.encode("utf-8")).hexdigest()[:10]
        nodes.append({"id": entity_id, "label": entity, "type": "Entity"})
        for filename in filenames:
            if filename in evidence_files:
                edges.append({
                    "source": entity_id,
                    "target": f"source:{filename}",
                    "relationship": "MENTIONED_IN",
                })
    return {"nodes": nodes, "edges": edges}


def _neo4j_context(source_names: list[str]) -> tuple[list[dict], str]:
    cypher = (
        "MATCH (a)-[r]-(b) "
        "WHERE a.source_file IN $source_files OR b.source_file IN $source_files "
        "RETURN labels(a)[0] AS source_type, coalesce(a.id, a.name, a.source_file) AS source, "
        "type(r) AS relationship, labels(b)[0] AS target_type, "
        "coalesce(b.id, b.name, b.source_file) AS target LIMIT 20"
    )
    if not source_names:
        return [], cypher
    try:
        from leaseguard.graph.neo4j_client import Neo4jClient

        with Neo4jClient() as client:
            rows = client.run(cypher, source_files=source_names)
        return [
            {key: _mask_pii(str(value))[0] for key, value in row.items()}
            for row in rows
        ], cypher
    except Exception as exc:
        logger.debug(f"Next AI Neo4j context unavailable: {exc}")
        return [], cypher


def _guard_question(question: str) -> Optional[str]:
    clean_question = question.strip()
    if not clean_question:
        return "Please enter a question."
    if len(clean_question) > 1200:
        return "The question is too long. Please keep it below 1,200 characters."
    lowered = clean_question.lower()
    if any(pattern in lowered for pattern in _BLOCKED_PATTERNS):
        return (
            "This request was blocked by the Next AI guardrail because it attempts to bypass "
            "controls, expose secrets, or perform a destructive action."
        )
    return None


def _fallback_dataset_questions(req: NextAIContextRequest, index: dict) -> list[str]:
    use_case = req.use_case or "this journey"
    source_terms = Counter()
    for chunk in index["chunks"][:80]:
        source_terms.update(_tokens(chunk["text"]))
    terms = [term for term, _ in source_terms.most_common(8)]
    source_names = [Path(source["filename"]).stem for source in index["sources"][:5]]
    entities = req.primary_entities[:3] or list(index["entity_sources"].keys())[:3]
    relationships = req.key_relationships[:3]
    topics = req.hint_topics[:3]
    focus = req.objective or use_case
    hints = [
        f"What evidence across the uploaded datasets best supports {focus}?",
        (
            "How are " + ", ".join(entities) +
            " connected across the loaded sources?"
            if entities else
            f"Which entities connect the uploaded sources for {use_case}?"
        ),
        (
            "Which records provide evidence for " + ", ".join(relationships) + "?"
            if relationships else
            f"What cross-source relationships are most relevant to {use_case}?"
        ),
        (
            "What patterns or exceptions involving " + ", ".join(topics) +
            " are supported by the database?"
            if topics else
            f"What risks, exceptions, or decision signals for {use_case} are supported by the database?"
        ),
        (
            "What combined insight can be derived from " + ", ".join(source_names[:3]) + "?"
            if source_names else
            f"What are the five strongest database-backed findings for {use_case}?"
        ),
    ]
    if terms:
        hints[4] = (
            "What does the uploaded evidence reveal about "
            + ", ".join(terms[:4]) + f" in relation to {use_case}?"
        )
    return hints[:5]


def _dataset_aware_questions(req: NextAIContextRequest, index: dict) -> tuple[list[str], str]:
    fallback = _fallback_dataset_questions(req, index)
    if not index["chunks"]:
        return fallback, "configured_fallback"
    source_context = []
    for source in index["sources"][:14]:
        source_chunks = [
            chunk["text"] for chunk in index["chunks"]
            if chunk["filename"] == source["filename"]
        ]
        source_context.append({
            "filename": source["filename"],
            "modality": source["modality"],
            "database_excerpt": _clean(" ".join(source_chunks[:2]), 1200),
        })
    common_terms = Counter()
    for chunk in index["chunks"][:120]:
        common_terms.update(_tokens(chunk["text"]))
    system = (
        "You generate the five most useful questions a user would ask an uploaded evidence "
        "database. Return STRICT JSON only with key questions containing exactly 5 unique strings. "
        "Every question must be tailored jointly to the supplied industry, use case, objective, "
        "configured entities/relationships, and actual uploaded dataset filenames/content. "
        "Questions must be answerable or meaningfully investigated from these datasets; do not "
        "invent unavailable facts. Prefer decision-useful joins, patterns, exceptions, provenance, "
        "comparisons, and cross-source relationships. Use concise natural language and end each "
        "question with a question mark."
    )
    user = json.dumps({
        "industry": req.industry,
        "use_case": req.use_case,
        "objective": req.objective,
        "primary_entities": req.primary_entities[:8],
        "key_relationships": req.key_relationships[:8],
        "configured_topics": req.hint_topics[:8],
        "uploaded_databases": source_context,
        "indexed_entities": list(index["entity_sources"].keys())[:30],
        "frequent_database_terms": [
            term for term, _ in common_terms.most_common(25)
        ],
    }, ensure_ascii=False)
    try:
        from leaseguard.llm.llm_client import LLMClient

        parsed = LLMClient().extract_json(system, user)
        values = parsed.get("questions") if isinstance(parsed, dict) else None
        questions = []
        for value in values or []:
            question = _clean(value, 220)
            if not question:
                continue
            if not question.endswith("?"):
                question += "?"
            if question.casefold() not in {item.casefold() for item in questions}:
                questions.append(question)
        if len(questions) >= 5:
            return questions[:5], "llm_dataset_synthesis"
    except Exception as exc:
        logger.warning(f"Next AI dataset-aware question synthesis failed: {exc}")
    return fallback, "dataset_heuristic_fallback"


def build_next_ai_status(req: NextAIContextRequest, db_client=None) -> dict:
    index = _build_index(req, db_client)
    hints, hint_source = _dataset_aware_questions(req, index)
    modality_counts = Counter(source["modality"] for source in index["sources"])
    source_count = len(index["sources"])
    use_case = req.use_case or "the configured use case"
    industry = req.industry or "the active domain"
    persona = req.persona or f"{use_case} evidence assistant"
    if source_count:
        welcome = (
            f"Hi — I am your {persona}, grounded in the active {industry} · {use_case} RAG database "
            f"({source_count} uploaded source(s), {len(index['chunks'])} indexed chunk(s), "
            f"and {len(index['entity_sources'])} extracted graph entities). Ask a question or "
            "try a journey-specific suggestion below."
        )
    else:
        welcome = (
            f"Hi — your {persona} is configured for {industry} · {use_case}, but the active upload "
            "session has no retrievable RAG evidence yet."
        )
    return {
        "industry": industry,
        "use_case": use_case,
        "tracking_id": req.tracking_id,
        "welcome": welcome,
        "hints": hints,
        "hint_source": hint_source,
        "source_count": source_count,
        "modality_counts": dict(modality_counts),
        "chunk_count": len(index["chunks"]),
        "entity_count": len(index["entity_sources"]),
        "pii_values_masked": index["pii_masked"],
        "vector_store": index["vector_store"],
        "agents": [
            {"name": "Session Ingestion", "status": "ready", "detail": f"{source_count} source(s)"},
            {"name": "Chunking & Indexing", "status": "ready", "detail": f"{len(index['chunks'])} chunk(s)"},
            {"name": "Embedding Search", "status": "ready", "detail": index["vector_store"]},
            {"name": "Entity & Relationship", "status": "ready", "detail": f"{len(index['entity_sources'])} entities"},
            {"name": "Graph-RAG & Cypher", "status": "ready", "detail": "Read-only graph expansion"},
            {"name": "Answer Synthesis", "status": "ready", "detail": "Evidence-restricted agent"},
            {"name": "PII Protection", "status": "ready", "detail": f"{index['pii_masked']} value(s) masked"},
            {"name": "Guardrail", "status": "ready", "detail": "Prompt and abuse controls active"},
        ],
    }


def _grounded_answer(
    req: NextAIAskRequest, evidence: list[dict], graph_rows: list[dict]
) -> str:
    context = [{
        "source": item["filename"],
        "similarity": item["score"],
        "excerpt": item["text"][:1400],
    } for item in evidence]
    system = (
        "You are Next AI, a privacy-preserving hybrid RAG and Graph-RAG assistant. Answer only "
        "from the supplied uploaded-evidence excerpts and graph rows. Never use outside facts. "
        "Preserve uncertainty and do not infer guilt, intent, identity, causation, legal status, "
        "or diagnosis. Cite exact filenames in square brackets after each material claim. PII "
        "has been masked and must remain masked. If evidence is insufficient, state that clearly."
    )
    user = json.dumps({
        "industry": req.industry,
        "use_case": req.use_case,
        "objective": req.objective,
        "question": req.question,
        "evidence": context,
        "graph_rows": graph_rows[:20],
        "recent_conversation": req.history[-6:],
    }, ensure_ascii=False)
    try:
        from leaseguard.llm.llm_client import LLMClient

        return _mask_pii(LLMClient().chat(system, user))[0]
    except Exception as exc:
        logger.warning(f"Next AI grounded synthesis failed: {exc}")
        top = evidence[0]
        return (
            f"The most relevant available evidence is in [{top['filename']}]: "
            f"{_clean(top['text'], 700)}"
        )


def _external_synthesis(req: NextAIAskRequest) -> str:
    system = (
        "You are Next AI. The user explicitly approved an ungrounded LLM synthesis because the "
        "active RAG database did not contain an answer. Give a concise general response, clearly "
        "label it as not derived from uploaded evidence, avoid personal data, and do not present "
        "speculation as fact."
    )
    try:
        from leaseguard.llm.llm_client import LLMClient

        answer = LLMClient().chat(
            system,
            json.dumps({
                "industry": req.industry,
                "use_case": req.use_case,
                "question": req.question,
            }, ensure_ascii=False),
        )
        return _mask_pii(answer)[0]
    except Exception as exc:
        logger.warning(f"Next AI external synthesis failed: {exc}")
        return "LLM synthesis is currently unavailable. Please try again later."


def answer_next_ai(req: NextAIAskRequest, db_client=None) -> dict:
    blocked = _guard_question(req.question)
    if blocked:
        return {
            "answer": blocked,
            "answer_available": False,
            "grounding": "blocked",
            "guardrail": {"blocked": True, "reason": blocked},
            "evidence": [],
            "graph_context": {"nodes": [], "edges": []},
            "graph_rows": [],
            "cypher_query": "",
            "retrieval": {},
        }

    index = _build_index(req, db_client)
    evidence = _semantic_search(index, req.question)
    summary_intent = any(word in req.question.lower() for word in ("summary", "summarize", "overview"))
    query_term_count = len(set(_tokens(req.question)))
    required_overlap = 1 if query_term_count <= 3 else 2
    relevant = [
        item for item in evidence
        if (
            item["score"] >= 0.08
            and item.get("term_overlap", 0) >= required_overlap
        ) or (summary_intent and item["score"] >= 0)
    ][:6]
    semantic_fallback = not bool(relevant) and bool(evidence)
    if semantic_fallback:
        # Continue with the closest uploaded chunks when lexical overlap is low.
        # The synthesis remains evidence-restricted and explicitly labelled.
        relevant = evidence[:6]
    local_graph = _local_graph_context(index, relevant, req.question)
    neo_rows, cypher = _neo4j_context([item["filename"] for item in relevant])

    if not relevant:
        if req.allow_llm_synthesis:
            answer = _external_synthesis(req)
            grounding = "llm_synthesis_not_rag_grounded"
        else:
            answer = (
                "This answer is not available in the active RAG database. The uploaded structured "
                "and unstructured evidence did not produce a sufficiently relevant semantic or "
                "graph match. You can explicitly ask Next AI to synthesize a general LLM response, "
                "which will be labelled as not grounded in uploaded data."
            )
            grounding = "not_available"
        return {
            "answer": answer,
            "answer_available": False,
            "can_synthesize_with_llm": not req.allow_llm_synthesis,
            "grounding": grounding,
            "guardrail": {"blocked": False, "pii_protected": True},
            "evidence": [],
            "graph_context": local_graph,
            "graph_rows": neo_rows,
            "cypher_query": cypher,
            "retrieval": {
                "semantic_candidates": len(evidence),
                "accepted_chunks": 0,
                "threshold": 0.08,
                "vector_store": index["vector_store"],
            },
        }

    answer = _grounded_answer(req, relevant, neo_rows)
    citations = list(dict.fromkeys(item["filename"] for item in relevant))
    return {
        "answer": answer,
        "answer_available": True,
        "can_synthesize_with_llm": False,
        "grounding": (
            "maximum_semantic_relevance_rag"
            if semantic_fallback else "hybrid_rag_graph_rag"
        ),
        "guardrail": {
            "blocked": False,
            "pii_protected": True,
            "pii_values_masked": index["pii_masked"],
        },
        "citations": citations,
        "evidence": [{
            "filename": item["filename"],
            "modality": item["modality"],
            "chunk_index": item["chunk_index"],
            "score": item["score"],
            "term_overlap": item.get("term_overlap", 0),
            "excerpt": _clean(item["text"], 500),
        } for item in relevant],
        "graph_context": local_graph,
        "graph_rows": neo_rows,
        "cypher_query": cypher,
        "retrieval": {
            "semantic_candidates": len(evidence),
            "accepted_chunks": len(relevant),
            "graph_nodes": len(local_graph["nodes"]),
            "graph_edges": len(local_graph["edges"]),
            "neo4j_rows": len(neo_rows),
            "threshold": 0.08,
            "semantic_fallback_used": semantic_fallback,
            "vector_store": index["vector_store"],
        },
    }

"""Requirement-specific cross-document discoveries generated from session sources."""
from __future__ import annotations

import csv
import io
import itertools
import json
import re
from pathlib import Path
from typing import Any, Optional

from loguru import logger
from pydantic import BaseModel, Field

from api.cypher_detections import _infer_industry


class CrossDocDiscoveryRequest(BaseModel):
    tabular_files: list[str] = Field(default_factory=list)
    document_files: list[str] = Field(default_factory=list)
    image_files: list[str] = Field(default_factory=list)
    audio_files: list[str] = Field(default_factory=list)
    text_files: list[str] = Field(default_factory=list)
    tracking_id: Optional[str] = None
    industry: Optional[str] = None
    use_case: Optional[str] = None
    requirement: str = ""
    objective: str = ""
    primary_entities: list[str] = Field(default_factory=list)
    key_relationships: list[str] = Field(default_factory=list)
    detection_focus: list[str] = Field(default_factory=list)
    page_title: str = ""
    page_subtitle: str = ""
    eyebrow: str = ""
    skip_llm: bool = False


def _modality(filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext in (".csv", ".tsv", ".xlsx", ".xls"):
        return "tabular"
    if ext in (".pdf", ".doc", ".docx"):
        return "document"
    if ext in (".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff"):
        return "image"
    if ext in (".wav", ".mp3", ".m4a", ".aac"):
        return "audio"
    return "text"


def _clean_text(value: Any, limit: int = 1800) -> str:
    text = re.sub(r"\s+", " ", str(value or "")).strip()
    return text[:limit]


def _extract_excerpt(filename: str, data: bytes) -> str:
    """Best-effort, bounded extraction for grounding the LLM prompt."""
    ext = Path(filename).suffix.lower()
    try:
        if ext in (".csv", ".tsv"):
            text = data.decode("utf-8", errors="replace")
            dialect = "\t" if ext == ".tsv" else ","
            rows = list(itertools.islice(csv.reader(io.StringIO(text), delimiter=dialect), 7))
            return _clean_text(" | ".join(" ; ".join(row[:10]) for row in rows))
        if ext in (".xlsx", ".xls"):
            import openpyxl

            book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            sheet = book.active
            rows = list(itertools.islice(sheet.iter_rows(values_only=True), 7))
            return _clean_text(" | ".join(" ; ".join(str(v or "") for v in row[:10]) for row in rows))
        if ext == ".pdf":
            try:
                from pypdf import PdfReader

                reader = PdfReader(io.BytesIO(data))
                return _clean_text(" ".join((p.extract_text() or "") for p in reader.pages[:3]))
            except Exception:
                import pdfplumber

                with pdfplumber.open(io.BytesIO(data)) as pdf:
                    return _clean_text(" ".join((p.extract_text() or "") for p in pdf.pages[:3]))
        if ext in (".docx", ".doc"):
            import docx

            document = docx.Document(io.BytesIO(data))
            return _clean_text(" ".join(p.text for p in document.paragraphs[:80]))
        if ext in (".txt", ".md", ".rtf", ".json", ".xml"):
            return _clean_text(data.decode("utf-8", errors="replace"))
    except Exception as exc:
        logger.debug(f"Cross-doc excerpt skipped for {filename}: {exc}")
    return ""


def _identifiers(text: str) -> list[str]:
    values = re.findall(r"\b[A-Z]{2,12}[-_/]?\d{2,12}\b|\b\d{8,16}\b", text.upper())
    return list(dict.fromkeys(values))[:30]


def _name_tokens(filename: str) -> set[str]:
    stop = {"csv", "pdf", "xlsx", "xls", "doc", "docx", "data", "dataset", "file"}
    return {
        token for token in re.findall(r"[a-z0-9]+", Path(filename).stem.lower())
        if len(token) > 2 and token not in stop
    }


def _request_files(req: CrossDocDiscoveryRequest) -> list[str]:
    return list(dict.fromkeys(
        req.tabular_files + req.document_files + req.image_files
        + req.audio_files + req.text_files
    ))


def _collect_sources(req: CrossDocDiscoveryRequest, db_client=None) -> list[dict]:
    sources: dict[str, dict] = {
        name: {
            "filename": name,
            "modality": _modality(name),
            "excerpt": "",
            "identifiers": [],
        }
        for name in _request_files(req)
    }
    if req.tracking_id and db_client:
        try:
            for item in db_client.get_files_for_session(req.tracking_id):
                filename = item.get("filename") or ""
                if not filename:
                    continue
                content = db_client.get_file_content(item["id"]) or {}
                excerpt = _extract_excerpt(filename, content.get("file_data") or b"")
                sources[filename] = {
                    "filename": filename,
                    "modality": _modality(filename),
                    "excerpt": excerpt,
                    "identifiers": _identifiers(excerpt),
                }
        except Exception as exc:
            logger.warning(f"Cross-doc session read failed: {exc}")
    return list(sources.values())


def _pair_score(left: dict, right: dict) -> tuple[float, list[str]]:
    shared_ids = sorted(set(left["identifiers"]) & set(right["identifiers"]))
    shared_tokens = _name_tokens(left["filename"]) & _name_tokens(right["filename"])
    modality_bonus = 4 if left["modality"] != right["modality"] else 0
    return len(shared_ids) * 12 + len(shared_tokens) * 2 + modality_bonus, shared_ids


def _candidate_findings(sources: list[dict], requirement: str) -> list[dict]:
    pairs = []
    for left, right in itertools.combinations(sources, 2):
        score, shared_ids = _pair_score(left, right)
        pairs.append((score, left, right, shared_ids))
    pairs.sort(key=lambda item: item[0], reverse=True)

    findings = []
    for index, (score, left, right, shared_ids) in enumerate(pairs[:5]):
        l_name, r_name = left["filename"], right["filename"]
        evidence = (
            f"Shared identifier(s): {', '.join(shared_ids[:5])}"
            if shared_ids
            else "Candidate selected from source names, modalities, and the active requirement; LLM validation follows."
        )
        findings.append({
            "id": f"candidate_{index + 1}",
            "title": f"Cross-source candidate: {Path(l_name).stem} and {Path(r_name).stem}",
            "summary": (
                f"Compare {left['modality']} evidence in {l_name} with "
                f"{right['modality']} evidence in {r_name} for the current requirement."
            ),
            "business_impact": requirement or "Potential evidence that is only visible when both sources are evaluated together.",
            "severity": "review",
            "confidence": min(0.95, 0.45 + score / 100),
            "sources": [
                {"filename": l_name, "modality": left["modality"]},
                {"filename": r_name, "modality": right["modality"]},
            ],
            "evidence": evidence,
            "graph_query": (
                "MATCH (a)-[r]-(b) WHERE a.source_file <> b.source_file "
                "RETURN a, type(r), b"
            ),
        })
    return findings


def build_crossdoc_payload(
    req: CrossDocDiscoveryRequest,
    db_client=None,
) -> dict:
    sources = _collect_sources(req, db_client)
    names_blob = " ".join(
        [s["filename"] for s in sources]
        + [req.requirement, req.industry or "", req.use_case or ""]
    ).lower()
    inferred_industry, inferred_use_case = _infer_industry(names_blob)
    industry = req.industry or inferred_industry
    use_case = req.use_case or inferred_use_case
    requirement = req.requirement or req.objective or use_case
    findings = _candidate_findings(sources, requirement)

    return {
        "source": "evidence_candidates",
        "industry": industry,
        "use_case": use_case,
        "requirement": requirement,
        "phase_label": f"Step 8 of 10 - {use_case} Cross-Document Discovery",
        "eyebrow": _clean_text(req.eyebrow, 160)
        or f"{industry} - Requirement-Specific Cross-Source Reasoning",
        "title": _clean_text(req.page_title, 180)
        or f"Findings Only Visible by Connecting Sources for {use_case}",
        "subtitle": _clean_text(req.page_subtitle, 350) or (
            f"Comparing {len(sources)} source(s) across the active {industry} requirement: "
            f"{requirement}. Findings and reasoning are generated from this session."
        ),
        "sources_analyzed": len(sources),
        "findings": findings,
    }


def enrich_crossdoc_with_llm(payload: dict, req: CrossDocDiscoveryRequest, db_client=None) -> dict:
    sources = _collect_sources(req, db_client)
    if len(sources) < 2:
        return payload

    source_context = [
        {
            "filename": source["filename"],
            "modality": source["modality"],
            "identifiers": source["identifiers"][:15],
            "excerpt": source["excerpt"][:1400],
        }
        for source in sources[:18]
    ]
    system = (
        "You are a cross-document knowledge-graph discovery planner. Return STRICT JSON only. "
        "Derive every finding specifically from the supplied industry, requirement, filenames, "
        "identifiers, and excerpts. Do not reuse a standard AML, healthcare, real-estate, or "
        "procurement checklist. Do not invent entities or facts absent from the evidence. "
        "A finding must require at least two named sources. If evidence is weak, label it as a "
        "candidate for validation rather than a confirmed fact. Tailor the page heading and every "
        "finding to the supplied domain configuration. Return keys: eyebrow, title, subtitle, "
        "findings. findings must contain 2..6 objects with keys: id, title, summary, "
        "business_impact, severity (critical|high|medium|review), confidence (0..1), "
        "sources (array of {filename, modality}), evidence, graph_query. Each graph_query must "
        "be a concise Cypher query tailored to that finding."
    )
    user = json.dumps({
        "industry": payload.get("industry"),
        "use_case": payload.get("use_case"),
        "requirement": payload.get("requirement"),
        "objective": req.objective,
        "primary_entities": req.primary_entities[:8],
        "key_relationships": req.key_relationships[:8],
        "detection_focus": req.detection_focus[:8],
        "configured_page": {
            "eyebrow": payload.get("eyebrow"),
            "title": payload.get("title"),
            "subtitle": payload.get("subtitle"),
        },
        "sources": source_context,
    }, ensure_ascii=False)

    try:
        from leaseguard.llm.llm_client import LLMClient

        parsed = LLMClient().extract_json(system, user)
        raw_findings = parsed.get("findings") if isinstance(parsed, dict) else None
        if not isinstance(raw_findings, list) or not raw_findings:
            return payload

        known = {source["filename"]: source["modality"] for source in sources}
        clean_findings = []
        for index, item in enumerate(raw_findings[:6]):
            if not isinstance(item, dict):
                continue
            clean_sources = []
            for source in item.get("sources") or []:
                if not isinstance(source, dict):
                    continue
                filename = str(source.get("filename") or "")
                if filename not in known:
                    continue
                clean_sources.append({
                    "filename": filename,
                    "modality": known[filename],
                })
            if len({s["filename"] for s in clean_sources}) < 2:
                continue
            severity = str(item.get("severity") or "review").lower()
            if severity not in {"critical", "high", "medium", "review"}:
                severity = "review"
            try:
                confidence = max(0.0, min(1.0, float(item.get("confidence", 0.5))))
            except (TypeError, ValueError):
                confidence = 0.5
            clean_findings.append({
                "id": str(item.get("id") or f"finding_{index + 1}")[:60],
                "title": _clean_text(item.get("title"), 140),
                "summary": _clean_text(item.get("summary"), 900),
                "business_impact": _clean_text(item.get("business_impact"), 700),
                "severity": severity,
                "confidence": confidence,
                "sources": clean_sources,
                "evidence": _clean_text(item.get("evidence"), 900),
                "graph_query": _clean_text(item.get("graph_query"), 800),
            })
        if not clean_findings:
            return payload
        payload["eyebrow"] = _clean_text(parsed.get("eyebrow"), 160) or payload["eyebrow"]
        payload["title"] = _clean_text(parsed.get("title"), 180) or payload["title"]
        payload["subtitle"] = _clean_text(parsed.get("subtitle"), 350) or payload["subtitle"]
        payload["findings"] = clean_findings
        payload["source"] = "llm"
    except Exception as exc:
        logger.warning(f"Cross-doc LLM enrichment failed: {exc}")
    return payload
